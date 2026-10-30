"""Integration coverage against the actual delivered v3.4 workbook, without caches."""
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from spatial_monitor.api import create_app
from spatial_monitor.io import load_dataset
from spatial_monitor.selection import layer_groups
from spatial_monitor.interpolation import interpolate_layer
from spatial_monitor.plotting import scene_figure

ROOT = Path(__file__).parents[1]
DEMO = ROOT / "examples/workbooks/monitoring_v34_demo.xlsx"
BLANK = ROOT / "examples/workbooks/monitoring_v34_blank.xlsx"


def changed(tmp_path, edit):
    book = load_workbook(DEMO)
    edit(book)
    target = tmp_path / "other_site.xlsx"
    book.save(target)
    book.close()
    return target


def test_real_workbook_without_formula_caches():
    book = load_workbook(DEMO, data_only=True)
    assert book['02_分析資料庫']['E2'].value is None
    book.close()
    data = load_dataset(DEMO)
    assert data.metadata['samples'] == 72
    assert len(data.points) == 12
    assert len(data.observations) == 288
    assert data.observations.valid.all()
    assert not data.issues.severity.eq('error').any()
    assert data.observations.source_row.tolist() == list(range(2, 290))
    assert data.observations.source_file.eq('monitoring_v34_demo.xlsx:02_分析資料庫').all()
    assert data.observations.site.eq('合成示範場址A（非實測）').all()
    assert data.geology.empty
    assert {'DO', 'DO_SAT', 'EC', 'ORP', 'pH', 'TOC', 'Zn'} == set(data.observations.parameter)
    assert data.observations.query("parameter == 'ORP'").value.min() < 0
    assert data.config['parameters']['pH']['transform'] == 'none'
    ec = data.observations.query("parameter == 'EC'")
    assert ec.value.iloc[0] == 430
    assert ec.comparison_temperature.eq(25).all()
    toc = data.observations.query("parameter == 'TOC'")
    assert set(toc.unit) == {'mg/kg', 'mg/L'}
    assert toc.query("medium == 'soil'").basis.eq('dry').all()


def test_soil_intervals_water_sampling_and_water_level_are_distinct():
    data = load_dataset(DEMO)
    soil = data.observations.query("medium == 'soil'")
    water = data.observations.query("medium == 'groundwater'")
    assert (soil.depth_bottom - soil.depth_top).eq(1).all()
    assert water.depth_top.eq(water.depth_bottom).all()
    assert set(water.depth_top) == {.5, 1.5, 2.5}
    assert len(data.water_levels) == 12
    assert data.water_levels.depth.eq(.2).all()
    assert data.water_levels.valid.all()
    group = data.observations.query("parameter == 'Zn'")
    _, layer = next(layer_groups(group))
    grid = interpolate_layer(layer, data.config['parameters']['Zn'], data.config['interpolation'])
    assert grid.available
    fig = scene_figure(group, data.geology, data.points, data.config['parameters']['Zn'])
    assert len(fig.data) > 0


def test_workbook_tables_formulas_and_blank_seed():
    demo, blank = load_workbook(DEMO), load_workbook(BLANK)
    assert demo.sheetnames == blank.sheetnames
    assert len(demo.sheetnames) == 8
    for name, count in [('Samples', 3), ('Results', 30), ('DB_03', 2)]:
        ws = next(s for s in demo if name in s.tables)
        assert sum(c.calculatedColumnFormula is not None for c in ws.tables[name].tableColumns) == count
    assert 'B73' in demo['01_樣品登錄']['L73'].value
    assert 'I289' in demo['02_分析資料庫']['E289'].value
    assert all(c.data_type == 'f' for c in demo['07_QAQC'][8])
    assert demo['01_樣品登錄']['C2'].is_date
    marker = [c.coordinate for row in demo['04_測項與清單'] for c in row if c.value == '下拉清單與自動選項']
    assert len(marker) == 1
    assert blank['04_測項與清單'][marker[0]].value == '下拉清單與自動選項'
    assert len(demo.defined_names) == 9
    assert load_dataset(BLANK).observations.empty
    demo.close()
    blank.close()


def test_missing_cache_stale_calculated_cells_and_blank_formatted_tail(tmp_path):
    def edit(w):
        w['02_分析資料庫']['E2'] = 'wrong cached point'
        w['02_分析資料庫']['G2'] = 999999
        w['01_樣品登錄']['L2'] = '地下水'
        w['02_分析資料庫']['A999'].number_format = '@'
    data = load_dataset(changed(tmp_path, edit))
    assert len(data.observations) == 288
    assert data.observations.iloc[0].value == 74
    assert data.observations.iloc[0].medium == 'soil'


@pytest.mark.parametrize('raw, qualifier', [('ND', 'ND'), ('＜0.050', '<'), ('≥10', '≥'), ('N/A', 'NA'), ('--', 'NA')])
def test_censored_source_is_retained_and_never_substituted(tmp_path, raw, qualifier):
    data = load_dataset(changed(tmp_path, lambda w: setattr(w['02_分析資料庫']['B2'], 'value', raw)))
    row = data.observations.iloc[0]
    assert row.raw_value == raw
    assert row.qualifier == qualifier
    assert not row.valid and pd.isna(row.value)
    if raw == '＜0.050':
        assert row.detection_limit == .05


def test_bad_links_duplicates_and_input_formulas_remain_visible(tmp_path):
    def edit(w):
        w['01_樣品登錄']['A3'] = w['01_樣品登錄']['A2'].value
        w['02_分析資料庫']['H4'] = 'not-a-sample'
        w['02_分析資料庫']['B5'] = '=1+1'
    data = load_dataset(changed(tmp_path, edit))
    assert len(data.observations) == 288
    assert {'unresolved_sample', 'invalid_identifier', 'invalid_source_cell'} <= set(data.issues.code)
    assert not data.observations.loc[data.observations.source_row.eq(5), 'valid'].any()


def test_groundwater_datum_is_not_silently_assumed(tmp_path):
    def edit(w):
        w['03_點位主檔']['L3'] = 'crs=EPSG:3826'
    data = load_dataset(changed(tmp_path, edit))
    assert 'unknown_water_datum' in data.issues.code.tolist()
    assert not data.water_levels.query("point_id == 'BH01-W'").valid.any()
    assert data.points.query("point_id == 'BH01-W'").surface_elevation.isna().all()


def test_month_only_date_is_not_completed_implicitly(tmp_path):
    data = load_dataset(changed(tmp_path, lambda w: setattr(w['01_樣品登錄']['C2'], 'value', '2026-01')))
    assert 'incomplete_date' in data.issues.code.tolist()
    assert not data.observations.iloc[0].valid
    assert pd.isna(data.observations.iloc[0].date)


def test_missing_ph_unit_and_disabled_master_are_not_assumed_valid(tmp_path):
    def edit(w):
        w['02_分析資料庫']['C4'] = None  # pH still requires its explicit dimensionless unit.
        w['04_測項與清單']['G6'] = '否'  # DO
    data = load_dataset(changed(tmp_path, edit))
    assert {'missing_unit', 'unresolved_parameter_master'} <= set(data.issues.code)
    assert not data.observations.query('source_row == 4').valid.any()
    assert not data.observations.query("parameter == 'DO'").valid.any()


def test_multiple_sites_use_their_own_names_and_coordinates(tmp_path):
    original = load_dataset(DEMO)
    def edit(w):
        w['00_使用說明']['C5'] = '另一個場址'
        for row in w['03_點位主檔'].iter_rows(min_row=2, max_row=13):
            row[4].value += 10000
    other = load_dataset(changed(tmp_path, edit))
    assert other.observations.site.eq('另一個場址').all()
    assert np.allclose(other.observations.x - original.observations.x, 10000)
    assert original.observations.site.eq('合成示範場址A（非實測）').all()
    pd.testing.assert_frame_equal(load_dataset(DEMO).observations, original.observations)


def test_master_defined_new_parameter_reaches_api_and_configuration(tmp_path):
    def edit(w):
        w['04_測項與清單']['B152'] = '自訂碳指標'
        w['02_分析資料庫']['A2'] = '自訂碳指標'
        w['04_測項與清單']['A152'] = '土壤'
        w['04_測項與清單']['F152'] = 'mg/kg'
        w['04_測項與清單']['G152'] = '是'
    path = changed(tmp_path, edit)
    data = load_dataset(path)
    assert data.observations.iloc[0].valid
    assert '自訂碳指標' in data.config['parameters']
    assert '自訂碳指標' not in load_dataset(DEMO).config['parameters']
    key = 'test-workbook-key-with-32-characters'
    with TestClient(create_app(api_keys={'test': key})) as client:
        result = client.post('/v1/validate/file', headers={'X-API-Key': key},
                             files={'file': ('other.xlsx', path.read_bytes())})
    assert result.status_code == 200, result.text
    assert result.json()['observations'] == 288
    assert result.json()['import']['format'] == 'monitoring-workbook-v3.4'
    assert result.json()['data'][0]['parameter'] == '自訂碳指標'
