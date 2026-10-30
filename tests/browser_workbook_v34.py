"""Verify actual workbook upload, plots, export, and two isolated browser sessions."""
import argparse
import json
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', default='http://127.0.0.1:8502')
    parser.add_argument('--browser', default='/usr/bin/chromium')
    parser.add_argument('--password-file', type=Path, default=ROOT / '.personal-site-password')
    parser.add_argument('--output', type=Path, default=ROOT / 'outputs/workbook-v34/browser')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    password = args.password_file.read_text().strip()
    demo = ROOT / 'examples/workbooks/monitoring_v34_demo.xlsx'
    other = args.output / 'site_b.xlsx'
    book = load_workbook(demo)
    book['00_使用說明']['C5'] = '合成測試場址B'
    for row in book['03_點位主檔'].iter_rows(min_row=2, max_row=13):
        row[4].value += 10000
    book.save(other)
    book.close()
    checks, errors = [], []

    def passed(name):
        checks.append(name)
        print('PASS', name, flush=True)

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=args.browser, headless=True,
            args=['--no-sandbox', '--disable-dev-shm-usage', '--enable-unsafe-swiftshader'])

        def login():
            context = browser.new_context(viewport={'width': 1440, 'height': 1050}, accept_downloads=True)
            page = context.new_page()
            page.set_default_timeout(60000)
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(args.url, wait_until='domcontentloaded')
            expect(page.get_by_role('button', name='進入工作台', exact=True)).to_be_visible(timeout=90000)
            page.get_by_role('textbox', name='測試密碼', exact=True).fill(password)
            page.get_by_role('button', name='進入工作台', exact=True).click()
            expect(page.get_by_role('tab')).to_have_count(6)
            return page

        def upload(page, path, site):
            uploader = page.get_by_test_id('stFileUploader').filter(has_text='Excel工作簿或多個CSV資料表')
            uploader.locator('input[type=file]').set_input_files(str(path))
            expect(page.get_by_text(f'已自動匯入 v3.4｜場址：{site}｜樣品：72筆', exact=True)).to_be_visible()
            expect(page.get_by_test_id('stMetricValue').nth(0)).to_have_text('288')
            expect(page.get_by_test_id('stMetricValue').nth(1)).to_have_text('288')
            expect(page.get_by_test_id('stMetricValue').nth(2)).to_have_text('12')
            expect(page.get_by_test_id('stException')).to_have_count(0)

        def select(page, label, value):
            page.get_by_role('combobox', name=label, exact=True).click()
            page.get_by_role('option', name=value, exact=True).click()

        a = login()
        a.get_by_text('v3.4範例、空白範本與座標設定', exact=True).click()
        with a.expect_download() as download:
            a.get_by_role('button', name='下載v3.4合成範例', exact=True).click()
        downloaded = args.output / 'downloaded_demo.xlsx'
        download.value.save_as(downloaded)
        assert downloaded.read_bytes() == demo.read_bytes()
        passed('Website serves the actual delivered v3.4 sample workbook')
        a.get_by_text('上傳Excel或CSV', exact=True).click()
        upload(a, downloaded, '合成示範場址A（非實測）')
        passed('User A uploads v3.4 without YAML: 72 samples, 288 valid results, 12 points')

        select(a, '參數', 'Zn')
        a.get_by_role('tab', name='分布與趨勢', exact=True).click()
        a.get_by_text('顯示2D內插估值', exact=True).click()
        a.wait_for_function("() => [...document.querySelectorAll('.js-plotly-plot')].some(p => p.offsetParent!==null && p.data?.some(t=>t.type==='heatmap'))")
        a.screenshot(path=str(args.output / 'v34-map.png'))
        passed('Native soil intervals produce a rendered 2D interpolation map')
        a.get_by_role('tab', name='3D實測展示', exact=True).click()
        expect(a.locator('.js-plotly-plot:visible canvas').first).to_be_visible()
        a.screenshot(path=str(args.output / 'v34-3d.png'))
        passed('Uploaded workbook renders an actual 3D measured view')

        a.get_by_role('tab', name='成果匯出', exact=True).click()
        a.get_by_role('button', name='產生成果壓縮檔', exact=True).click()
        expect(a.get_by_role('button', name='下載完整成果ZIP', exact=True)).to_be_visible(timeout=180000)
        with a.expect_download() as download:
            a.get_by_role('button', name='下載完整成果ZIP', exact=True).click()
        archive = args.output / 'native-results.zip'
        download.value.save_as(archive)
        with zipfile.ZipFile(archive) as z:
            assert z.testzip() is None
            assert json.loads(z.read('manifest.json'))['observations'] == 288
            assert json.loads(z.read('import.json'))['site'] == '合成示範場址A（非實測）'
        passed('Full native-workbook results download as a valid ZIP with source metadata')

        b = login()
        b.get_by_text('上傳Excel或CSV', exact=True).click()
        upload(b, other, '合成測試場址B')
        expect(a.get_by_text('已自動匯入 v3.4｜場址：合成示範場址A（非實測）｜樣品：72筆', exact=True)).to_be_visible()
        expect(a.get_by_role('button', name='下載完整成果ZIP', exact=True)).to_be_visible()
        b.get_by_role('tab', name='成果匯出', exact=True).click()
        expect(b.get_by_role('button', name='下載完整成果ZIP', exact=True)).to_have_count(0)
        passed('Two independently logged-in browser sessions keep their sites and outputs separate')

        # Streamlit multi-file uploader appends files: remove A before uploading B.
        a.get_by_role('button', name='Remove ' + downloaded.name, exact=True).click()
        upload(a, other, '合成測試場址B')
        expect(a.get_by_role('button', name='下載完整成果ZIP', exact=True)).to_have_count(0)
        expect(a.get_by_role('combobox', name='參數', exact=True)).to_have_value('全部')
        passed('Changing sites in one session clears old filters and exported results')
        assert not errors, errors
        browser.close()
    (args.output / 'report.json').write_text(json.dumps({'url': args.url, 'checks': checks,
        'passed': len(checks), 'page_errors': errors}, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
