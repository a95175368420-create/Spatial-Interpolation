"""Native v3.4 import from source cells, independent of Excel calculation caches.

The eight-sheet workbook is one site. Formula columns are evidence only, never
measurement inputs. Extra measurement metadata may be recorded as key=value
pairs in 備註, without changing the workbook's table or formula structure.
"""
from collections import Counter
from datetime import date, datetime
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


SHEETS = ("00_使用說明", "01_樣品登錄", "02_分析資料庫", "03_點位主檔", "04_測項與清單")
ALIASES = {"鋅": "Zn", "酸鹼值": "pH", "酸鹼度": "pH", "電導度": "EC",
           "溶氧": "DO", "溶氧量": "DO", "溶氧飽和度": "DO_SAT", "氧化還原電位": "ORP",
           "總有機碳": "TOC", "水溫": "temperature", "溫度": "temperature", "濁度": "turbidity"}
NOTE_KEYS = {"basis", "method", "temperature", "temperature_compensated", "reference_temperature",
             "reference_electrode", "calibration", "crs", "elevation_datum", "wellhead_offset"}


def notes(value):
    """Only explicit, documented metadata keys are interpreted; other prose stays raw."""
    return {key: value.strip() for key, value in
            re.findall(r"(?:^|[;；\n])\s*([a-z_]+)\s*=\s*([^;；\n]+)", str(value))
            if key in NOTE_KEYS}


def read_workbook(path, config, issues):
    from .io import FIELDS, MEDIUM, _unit

    def finding(table, row, code, message, severity="warning"):
        issues.append(dict(severity=severity, table=table, row=row, code=code, message=message))

    def text(value):
        return str(value).strip() if value is not None else ""

    def records(book, sheet, table, headers, logical):
        ws = book[sheet]
        if table not in ws.tables:
            raise ValueError(f"{sheet}: 缺少資料表 {table}，請保留 v3.4 原生表格")
        left, top, right, bottom = range_boundaries(ws.tables[table].ref)
        positions = {text(ws.cell(top, col).value): col for col in range(left, right + 1)}
        missing = set(headers) - set(positions)
        if missing:
            raise ValueError(f"{sheet}: 缺少原始輸入欄位 {', '.join(sorted(missing))}")
        # Source cells determine record counts, not formatting or prefilled formulas.
        candidates = sorted({r for (r, c), cell in ws._cells.items()
                             if top < r and c in {positions[h] for h in headers}
                             and cell.value is not None and (r <= bottom or logical in {"points", "samples", "analyses"})})
        rows = []
        for row in candidates:
            values = {h: ws.cell(row, positions[h]).value for h in headers}
            if not any(text(v) for v in values.values()):
                continue
            valid = True
            for header in headers:
                cell = ws.cell(row, positions[header])
                if cell.data_type in ("f", "e"):
                    valid = False
                    finding(logical, row, "invalid_source_cell", f"{sheet}!{cell.coordinate}: 輸入欄不可為公式或Excel錯誤", "error")
            if row > bottom:
                finding(logical, row, "outside_table", f"{sheet} 第{row}列在表格範圍外；已保留，請延伸Excel表格")
            rows.append({**{k: v if v is not None else "" for k, v in values.items()},
                         "source_file": f"{path.name}:{sheet}", "source_row": row, "_source_valid": valid})
        return rows

    book = load_workbook(path, data_only=False)
    try:
        missing = set(SHEETS) - set(book.sheetnames)
        if missing:
            raise ValueError(f"v3.4 工作簿缺少工作表：{', '.join(sorted(missing))}")
        site_cell = book[SHEETS[0]]["C5"]
        site = text(site_cell.value)
        if not site or site_cell.data_type in ("f", "e"):
            raise ValueError("00_使用說明!C5 請填入此工作簿的場址名稱（文字）")
        points = records(book, SHEETS[3], "DB_03", ["監測點", "點位類型", "預設介質", "井深(m)",
            "TWD97(X)", "TWD97(Y)", "絕對高程(m)", "篩管頂深(m)", "篩管底深(m)", "井管直徑(m)", "啟用", "備註"], "points")
        samples = records(book, SHEETS[1], "Samples", ["樣品編號", "監測點", "採樣日期", "採樣說明",
            "採樣深度(m)", "深度起(m)", "深度迄(m)", "採樣方法", "採樣單位", "檢測單位", "備註"], "samples")
        results = records(book, SHEETS[2], "Results", ["分析項目", "原始值", "結果單位", "備註", "關聯鍵"], "analyses")
        master = records(book, SHEETS[4], "DB_04", ["介質", "分析項目", "別名/英文", "測項類別", "預設分析方法", "標準單位", "啟用"], "parameters")
    finally:
        book.close()

    point_counts = Counter(text(p["監測點"]) for p in points)
    for p in points:
        meta = notes(p["備註"])
        crs = config.get("crs") or meta.get("crs") or "EPSG:3826"
        if not config.get("crs") and not meta.get("crs"):
            finding("points", p["source_row"], "assumed_crs", "TWD97採預設121度分帶 EPSG:3826；119度分帶請在匯入設定選擇 EPSG:3825")
        p.update(point_id=text(p["監測點"]), site=site, x=p["TWD97(X)"], y=p["TWD97(Y)"], crs=crs,
                 surface_elevation=p["絕對高程(m)"] if meta.get("elevation_datum") == "ground" else "")
        if p["絕對高程(m)"] != "" and meta.get("elevation_datum") != "ground":
            finding("points", p["source_row"], "unspecified_elevation_datum", "絕對高程原值保留；地表基準未確認，展示使用相對深度。備註可標記 elevation_datum=ground")
        if p["啟用"] != "是":
            p["_source_valid"] = False
            finding("points", p["source_row"], "inactive_point", "點位未啟用；保留來源但不參與推估", "error")
    unique_points = {p["point_id"]: p for p in points if point_counts[p["point_id"]] == 1}
    water = []
    for s in samples:
        p = unique_points.get(text(s["監測點"]), {})
        raw_medium = text(p.get("預設介質", ""))
        medium = MEDIUM.get(raw_medium, raw_medium.lower())
        top, bottom = ((s["深度起(m)"], s["深度迄(m)"]) if medium == "soil"
                       else (s["採樣深度(m)"], s["採樣深度(m)"]))
        s.update(sample_id=text(s["樣品編號"]), point_id=text(s["監測點"]), date=s["採樣日期"],
                 medium=medium, depth_top=top, depth_bottom=bottom, site=site, batch=text(s["採樣說明"]))
        date_value = s["採樣日期"]
        if not (isinstance(date_value, (datetime, date)) or (isinstance(date_value, str) and re.fullmatch(
                r"(?:\d{8}|\d{4}[-/]\d{1,2}[-/]\d{1,2})(?:[ T]\d{2}:\d{2}(?::\d{2})?)?", text(date_value)))):
            s["_source_valid"] = False
            s["date"] = ""
            finding("samples", s["source_row"], "incomplete_date", "採樣日期需要真實Excel日期或完整年月日；不以月資料或日期序號補造日期", "error")
        if medium == "groundwater" and s["深度起(m)"] != "":
            # F is depth below wellhead. Only a documented offset can locate it
            # against ground level. Unknown datums remain visible in the QA table.
            offset = notes(p.get("備註", "")).get("wellhead_offset")
            depth = pd.to_numeric(s["深度起(m)"], errors="coerce")
            shift = pd.to_numeric(offset, errors="coerce")
            known = pd.notna(shift) and pd.notna(depth)
            water.append(dict(point_id=s["point_id"], date=s["date"], depth=depth - shift if known else depth,
                              raw_wellhead_depth=s["深度起(m)"], wellhead_offset=offset or "",
                              depth_reference="ground" if known else "wellhead", _source_valid=known and s["_source_valid"],
                              source_file=s["source_file"], source_row=s["source_row"]))
            if not known:
                finding("water_levels", s["source_row"], "unknown_water_datum", "井口水位已保留，未提供 wellhead_offset（井口高於地面公尺數），不套疊為地表水位")

    sample_counts = Counter(s["sample_id"] for s in samples)
    unique_samples = {s["sample_id"]: s for s in samples if sample_counts[s["sample_id"]] == 1}
    counts = Counter((text(r["介質"]), text(r["分析項目"])) for r in master)
    definitions = {}
    for m in master:
        raw_medium, raw_name = text(m["介質"]), text(m["分析項目"])
        key = (raw_medium, raw_name)
        if counts[key] != 1 or m["啟用"] != "是" or not m["_source_valid"] or not text(m["標準單位"]):
            continue
        medium = MEDIUM.get(raw_medium, raw_medium.lower())
        parameter = ALIASES.get(raw_name, raw_name)
        unit = _unit(m["標準單位"])
        if parameter == "DO" and unit == "%":
            parameter = "DO_SAT"
        definitions[key] = (parameter, m)
        if parameter not in config["parameters"]:
            config["parameters"][parameter] = dict(label=raw_name, unit=unit, media=[], unit_by_medium={},
                data_type="continuous", transform="none", colorscale="Viridis", valid_range=None, thresholds=[])
        definition = config["parameters"][parameter]
        if medium not in definition.setdefault("media", []):
            definition["media"].append(medium)
        # Known parameters retain validated target units; a new medium gets the
        # workbook's declared unit. Actual conversions still use the shared validator.
        definition.setdefault("unit_by_medium", {}).setdefault(medium, definition["unit"] if
            parameter in {"pH", "EC", "DO", "DO_SAT", "ORP", "temperature", "turbidity"} else unit)

    for r in results:
        if not text(r["結果單位"]):
            r["_source_valid"] = False
            finding("analyses", r["source_row"], "missing_unit", "結果單位不得留白；無因次測項請填無因次", "error")
        s = unique_samples.get(text(r["關聯鍵"]), {})
        p = unique_points.get(s.get("point_id"), {})
        key = (text(p.get("預設介質", "")), text(r["分析項目"]))
        registered = definitions.get(key)
        parameter, m = registered if registered else (ALIASES.get(key[1], key[1]), {})
        if parameter == "DO" and _unit(r["結果單位"]) == "%":
            parameter = "DO_SAT"
        metadata = {**notes(s.get("備註", "")), **notes(r["備註"])}
        r.update(sample_id=text(r["關聯鍵"]), parameter=parameter, value=r["原始值"], unit=text(r["結果單位"]),
                 method=metadata.get("method", text(m.get("預設分析方法", ""))),
                 **{k: metadata.get(k, "") for k in NOTE_KEYS - {"method", "crs", "elevation_datum", "wellhead_offset"}})
        if not registered:
            r["_source_valid"] = False
            finding("analyses", r["source_row"], "unresolved_parameter_master", f"04測項主檔未有唯一啟用的介質／測項：{key}", "error")

    tables = {}
    for name, rows in {"samples": samples, "points": points, "analyses": results,
                       "water_levels": water, "geology": []}.items():
        frame = pd.DataFrame(rows)
        for col in [*FIELDS[name], "source_file", "source_row"]:
            if col not in frame:
                frame[col] = ""
        tables[name] = frame.fillna("")
    metadata = dict(format="monitoring-workbook-v3.4", site=site, samples=len(samples),
                    source=path.name, depth_rule="土壤F/G區間；水樣E單點深度；地下水F為井口水位",
                    standards="保留工作簿標準；平台不自動採用未確認適用性的法規門檻",
                    geology="v3.4沒有地質區間資料表；不自動建立岩性或地層")
    finding("workbook", None, "native_v34", "已辨識v3.4：以02關聯鍵連接01樣品及03點位，重新解析原始值；不使用公式快取", "info")
    return tables, metadata
